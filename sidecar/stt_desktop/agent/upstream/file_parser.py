from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from fastapi import HTTPException, status
from openpyxl import load_workbook
from stt_desktop.agent.upstream.official_workbooks import parse_official_workbook
from stt_desktop.agent.upstream.workbook_presentation import EXAMPLE_SHEET


SUPPORTED_EXTENSIONS = {".txt", ".csv", ".xlsx", ".doc", ".docx"}


def parse_attachment(filename: str, data: bytes, max_file_mb: int) -> str:
    return parse_attachment_content(filename, data, max_file_mb)["text"]


def parse_attachment_content(filename: str, data: bytes, max_file_mb: int) -> dict:
    max_bytes = max_file_mb * 1024 * 1024
    if len(data) > max_bytes:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File is larger than {max_file_mb} MB",
        )
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only .txt, .csv, .xlsx, .doc and .docx files are supported",
        )
    if suffix == ".xlsx":
        official = parse_official_workbook(data)
        if official is not None:
            return {"text": f"官方 {official['scene']} 模板 v{official['version']}，完整数据由官方模板校验器处理。", "data": official}
        return _parse_xlsx_content(data)
    if suffix == ".docx":
        parsed = _parse_docx(data)
        return {"text": parsed, "data": {"kind": "document", "lines": parsed.splitlines()}}
    if suffix == ".doc":
        parsed = _parse_legacy_doc(data)
        return {"text": parsed, "data": {"kind": "document", "lines": parsed.splitlines()}}
    text = _decode_text(data)
    if suffix == ".csv":
        return _parse_csv_content(text)
    parsed = text
    return {"text": parsed, "data": {"kind": "text", "lines": parsed.splitlines()}}


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_csv(text: str) -> str:
    return _parse_csv_content(text)["text"]


def _parse_csv_content(text: str) -> dict:
    sample = io.StringIO(text)
    rows = []
    reader = csv.reader(sample)
    structured_rows = []
    for index, row in enumerate(reader):
        if index >= 200:
            raise HTTPException(400, "CSV超过200行，请拆分后导入；不会截断数据。")
        rows.append("\t".join(row))
        structured_rows.append({"row": index + 1, "values": row})
    parsed = "\n".join(rows)
    return {"text": parsed, "data": {"kind": "table", "sheets": [{"name": "CSV", "rows": structured_rows}]}}


def _parse_xlsx(data: bytes) -> str:
    return _parse_xlsx_content(data)["text"]


def _parse_xlsx_content(data: bytes) -> dict:
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    try:
        visible = [sheet for sheet in workbook if sheet.sheet_state == "visible" and sheet.title != EXAMPLE_SHEET]
        if len(visible) > 20 or sum(sheet.max_row * sheet.max_column for sheet in visible) > 100000:
            raise HTTPException(400, "Excel 内容过多，请拆分为不超过 20 个工作表、10 万个单元格的文件。")
        parts: list[str] = []
        sheets: list[dict] = []
        for sheet in visible:
            if sheet.max_row > 2000 or sheet.max_column > 64:
                raise HTTPException(400, f"{sheet.title}：最多支持 2000 行、64 列，请拆分后上传。")
            # Expand only explicit merges; ordinary blank cells must remain blank.
            merged_values = {}
            for area in sheet.merged_cells.ranges:
                value = sheet.cell(area.min_row, area.min_col).value
                for row in range(area.min_row, area.max_row + 1):
                    for column in range(area.min_col, area.max_col + 1):
                        merged_values[row, column] = value
            parts.append(f"# {sheet.title}")
            rows = []
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                values = [merged_values.get((row_index, cell.column), cell.value) for cell in row]
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                values = ["" if value is None else str(value) for value in values]
                parts.append("\t".join(values))
                rows.append({"row": row_index, "values": values})
            sheets.append({"name": sheet.title, "rows": rows})
        return {"text": "\n".join(parts), "data": {"kind": "workbook", "sheets": sheets}}
    finally:
        workbook.close()


def _parse_docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            document_xml = archive.read("word/document.xml")
    except (KeyError, zipfile.BadZipFile) as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Could not read .docx document text",
        ) from exc

    try:
        root = ElementTree.fromstring(document_xml)
    except ElementTree.ParseError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Could not parse .docx document text",
        ) from exc

    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    lines: list[str] = []
    for paragraph in root.findall(".//w:p", namespace):
        pieces: list[str] = []
        for node in paragraph.iter():
            tag = node.tag.rsplit("}", 1)[-1]
            if tag == "t" and node.text:
                pieces.append(node.text)
            elif tag == "tab":
                pieces.append("\t")
            elif tag == "br":
                pieces.append("\n")
        line = "".join(pieces).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def _parse_legacy_doc(data: bytes) -> str:
    text = _extract_legacy_doc_text(data)
    if not text.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Could not extract readable text from .doc file. Please convert it to .docx if this keeps happening.",
        )
    return text


def _extract_legacy_doc_text(data: bytes) -> str:
    candidates: list[str] = []
    for encoding in ("utf-16le", "gb18030", "utf-8", "latin1"):
        decoded = data.decode(encoding, errors="ignore")
        candidates.extend(_readable_text_runs(decoded))

    seen: set[str] = set()
    useful: list[str] = []
    for candidate in candidates:
        cleaned = _clean_legacy_doc_line(candidate)
        if len(cleaned) < 2 or cleaned in seen:
            continue
        if not _looks_like_human_text(cleaned):
            continue
        seen.add(cleaned)
        useful.append(cleaned)
        if len("\n".join(useful)) >= 20000:
            break
    text = "\n".join(useful)
    marker_index = text.find("WpsCustomData")
    if marker_index >= 0:
        candidate = text[marker_index + len("WpsCustomData") :].strip()
        if len(re.findall(r"[\u4e00-\u9fff]", candidate)) >= 20:
            return candidate
    return text


def _readable_text_runs(text: str) -> list[str]:
    pattern = r"[\u4e00-\u9fffA-Za-z0-9（）()，,。.:：；;、/\-—_＋+%#&·\s]{3,}"
    return re.findall(pattern, text)


def _clean_legacy_doc_line(value: str) -> str:
    cleaned = re.sub(r"\s+", " ", value.replace("\x00", " ")).strip()
    cleaned = re.sub(r"^[^\u4e00-\u9fffA-Za-z0-9]+", "", cleaned)
    cleaned = re.sub(r"[^\u4e00-\u9fffA-Za-z0-9）)]$", "", cleaned)
    return cleaned.strip()


def _looks_like_human_text(value: str) -> bool:
    chinese_count = len(re.findall(r"[\u4e00-\u9fff]", value))
    alpha_count = len(re.findall(r"[A-Za-z]", value))
    digit_count = len(re.findall(r"\d", value))
    return chinese_count >= 2 or (alpha_count + digit_count >= 6 and len(value) <= 240)
