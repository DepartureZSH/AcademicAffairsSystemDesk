import io
import json
import pytest
from openpyxl import load_workbook
from stt_desktop.storage import ProjectWorkspace, ProjectError
from stt_desktop.agent.service import AgentService
from stt_desktop.agent.repository import LocalAgentRepository
from stt_desktop.agent.upstream.tools import execute_agent_action_items
from stt_desktop.agent.upstream.official_workbooks import (
    workbook_bytes,
    planning_collection_workbook_bytes,
)


@pytest.fixture
def agent(tmp_path):
    workspace = ProjectWorkspace(tmp_path / "workspace")
    with workspace.create_project("AI 测试") as project:
        yield AgentService(project, workspace)


def proposal(target, **payload):
    return {
        "target": target,
        "operation": "create",
        "payload": payload,
        "human_summary": "测试新增",
    }


@pytest.mark.parametrize("scene", ["timetable", "rooms", "school", "planning", "constraints"])
def test_every_scene_builds_its_real_provider_prompt(agent, scene):
    turn = agent.start(scene, "请检查当前资料")
    assert turn["providerRequest"]["tools"]
    assert turn["providerRequest"]["messages"][0]["role"] == "system"


def test_room_dependencies_review_execute_and_retry(agent):
    before = agent.project.revision
    action = agent.compile(
        "rooms",
        [
            proposal("room", name="物理实验室", room_type_name="实验室", capacity=40),
            proposal("room_type", name="实验室"),
        ],
    )
    assert agent.project.revision == before
    assert not agent.project.list_entities("room")
    assert action["items"][0]["target"] == "room_type"
    result = agent.confirm(action["id"])
    assert result["status"] == "executed"
    assert agent.project.revision == before + 1
    assert agent.confirm(action["id"])["status"] == "executed"
    assert len(agent.project.list_entities("room")) == 1
    assert (
        agent.project.connection.execute("SELECT count(*) FROM backup_records").fetchone()[0] == 1
    )


def test_reject_stale_and_atomic_failure(agent):
    action = agent.compile(
        "school", [proposal("teacher", name="张三"), proposal("teacher", name="李四")]
    )
    # Exercise a late failing item after the first SQL write, not just compile rejection.
    broken = json.loads(json.dumps(action))
    broken["items"][1]["after_data"]["unknown_field"] = 1
    with pytest.raises(ProjectError):
        execute_agent_action_items(
            LocalAgentRepository(agent.project, broken), broken, broken["items"]
        )
    assert not agent.project.list_entities("teacher")
    assert agent.get("action", action["id"])["status"] == "pending_confirmation"
    agent.project.save_entity("room", {"name": "已有教室"}, agent.project.revision)
    with pytest.raises(ProjectError, match="项目已变化"):
        agent.confirm(action["id"])
    assert agent.reject(action["id"])["status"] == "rejected"
    with pytest.raises(ProjectError):
        agent.confirm(action["id"])


def tool(name, args, identifier="call1"):
    return {
        "role": "assistant",
        "content": None,
        "tool_calls": [
            {
                "id": identifier,
                "type": "function",
                "function": {"name": name, "arguments": json.dumps(args)},
            }
        ],
    }


def test_provider_exchange_query_propose_finalize(agent):
    turn = agent.start("rooms", "创建一个实验室类型")
    assert turn["providerRequest"]["tools"][0]["function"]["name"] == "query_project_data"
    turn = agent.step(
        turn["turnId"], turn["step"], tool("query_project_data", {"dataset": "school_data"})
    )
    assert json.loads(turn["providerRequest"]["messages"][-1]["content"])["ok"]
    turn = agent.step(
        turn["turnId"],
        turn["step"],
        tool("propose_agent_action", proposal("room_type", name="实验室")),
    )
    result = agent.step(
        turn["turnId"], turn["step"], {"role": "assistant", "content": "请审核新增实验室类型。"}
    )
    assert result["done"] and len(result["actions"]) == 1
    assert not agent.project.list_entities("room_type")
    assert (
        AgentService(agent.project, agent.workspace).session("rooms")["actions"]
        == result["actions"]
    )
    agent.confirm(result["actions"][0]["id"])
    assert len(agent.project.list_entities("room_type")) == 1


def test_unknown_tools_and_cross_scene_rejected(agent):
    turn = agent.start("school", "查询资料")
    for name, args in [
        ("execute_sql", {"sql": "DELETE FROM teachers"}),
        ("query_project_data", {"dataset": "constraints"}),
        ("propose_agent_action", proposal("room", name="非法")),
    ]:
        turn = agent.step(turn["turnId"], turn["step"], tool(name, args))
        assert not json.loads(turn["providerRequest"]["messages"][-1]["content"])["ok"]
    assert not agent.project.list_entities("room")


@pytest.mark.parametrize("scene", ["rooms", "school", "planning"])
def test_original_excel_templates(scene):
    data = planning_collection_workbook_bytes() if scene == "planning" else workbook_bytes(scene)
    book = load_workbook(io.BytesIO(data))
    assert len(book.sheetnames) >= 2
    assert len(data) > 4000


def test_subject_minutes(agent):
    action = agent.compile(
        "school", [proposal("subject", name="数学", default_duration_minutes=40)]
    )
    agent.confirm(action["id"])
    subject = agent.project.list_entities("subject")[0]
    assert subject["default_duration_minutes"] == 40 and subject["default_duration_slots"] == 1


def test_plan_and_lesson_execution(tmp_path):
    from test_scheduling import seed_project

    project, task, lessons = seed_project(tmp_path, slot_count=3)
    with project:
        agent = AgentService(project, ProjectWorkspace(tmp_path / "workspace"))
        teacher = project.list_entities("teacher")[0]
        plan, _ = project.save_entity(
            "course_plan",
            {
                "term_id": task["term_id"],
                "homeroom_id": task["homeroom_id"],
                "subject_id": task["subject_id"],
                "weekly_slots": len(lessons),
            },
            project.revision,
        )
        project.save_entity(
            "teaching_task", {"id": task["id"], "course_plan_id": plan["id"]}, project.revision
        )
        # Update a real task and preserve existing lesson identities/preferences.
        action = agent.compile(
            "planning",
            [
                {
                    "target": "task",
                    "operation": "update",
                    "payload": {
                        "id": task["id"],
                        "homeroom_id": task["homeroom_id"],
                        "subject_id": task["subject_id"],
                        "primary_teacher_id": teacher["id"],
                        "fixed_room_id": None,
                        "required_room_type": "__no_room__",
                    },
                }
            ],
        )
        assert any(item["target"] == "task" for item in action["items"])
        agent.confirm(action["id"])
        assert {lesson["id"] for lesson in project.list_entities("task_lesson")} == {
            lesson["id"] for lesson in lessons
        }


def test_timetable_create(agent):
    action = agent.compile(
        "timetable",
        [
            proposal(
                "timetable_template",
                name="上午作息",
                day_count=5,
                periods=[
                    {
                        "weekday": 1,
                        "period_index": 1,
                        "label": "第一节",
                        "start_time_minutes": 480,
                        "end_time_minutes": 520,
                        "active": True,
                    }
                ],
            )
        ],
    )
    assert not agent.project.list_entities("bell_schedule")
    agent.confirm(action["id"])
    assert agent.project.list_entities("bell_schedule")[0]["name"] == "上午作息"
    assert len(agent.project.list_entities("time_slot")) == 1


def test_api_scope_confirmation_and_template_download(tmp_path):
    from test_sidecar_api import client, headers

    with client(tmp_path) as api:
        created = api.post("/v1/projects", headers=headers(), json={"name": "AI 集成测试"})
        assert created.status_code == 201
        pid = created.json()["project"]["id"]
        assert (
            api.get(
                "/v1/ai/session",
                headers=headers(),
                params={"project_id": "wrong", "scene": "rooms"},
            ).status_code
            == 400
        )
        start = api.post(
            "/v1/ai/turns",
            headers=headers(),
            json={"project_id": pid, "scene": "rooms", "content": "创建实验室类型"},
        ).json()
        response = api.post(
            f"/v1/ai/turns/{start['turnId']}/step",
            headers=headers(),
            json={
                "project_id": pid,
                "step": 0,
                "message": tool("propose_agent_action", proposal("room_type", name="实验室")),
            },
        )
        assert response.status_code == 200
        done = api.post(
            f"/v1/ai/turns/{start['turnId']}/step",
            headers=headers(),
            json={
                "project_id": pid,
                "step": 1,
                "message": {"role": "assistant", "content": "请审核。"},
            },
        ).json()
        action = done["actions"][0]
        assert (
            api.post(
                f"/v1/ai/actions/{action['id']}/confirm",
                headers=headers(),
                json={"project_id": "wrong"},
            ).status_code
            == 400
        )
        result = api.post(
            f"/v1/ai/actions/{action['id']}/confirm", headers=headers(), json={"project_id": pid}
        )
        assert result.status_code == 200 and result.json()["status"] == "executed"
        assert api.post("/v1/ai/workbook", json={"scene": "rooms"}).status_code == 401
        download = api.post("/v1/ai/workbook", headers=headers(), json={"scene": "rooms"})
        assert download.status_code == 200 and download.json()["data"].startswith("UEs")
        clear = api.post(
            "/v1/ai/clear", headers=headers(), json={"project_id": pid, "scene": "rooms"}
        )
        assert clear.status_code == 200
        session = api.get(
            "/v1/ai/session", headers=headers(), params={"project_id": pid, "scene": "rooms"}
        ).json()
        assert session["messages"] == [] and len(session["actions"]) == 1


def test_pending_questions_block_proposals(agent):
    turn = agent.start("rooms", "新增教室")
    turn = agent.step(
        turn["turnId"], 0, tool("propose_agent_action", proposal("room", name="未确认教室"))
    )
    done = agent.step(
        turn["turnId"],
        1,
        {"role": "assistant", "content": '```user_questions\n["是否新增？"]\n```'},
    )
    assert done["actions"] == [] and not agent.project.list_entities("room")
