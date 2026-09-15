"""Versioned, bounded workbook contracts for predictable, approval-only imports."""
from __future__ import annotations

import io
import re
from datetime import time
from zipfile import ZipFile, BadZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from stt_desktop.agent.upstream.compiler import ActionCompilationError
from stt_desktop.agent.upstream.workbook_layouts import LAYOUTS, read_layout
from stt_desktop.agent.upstream.workbook_presentation import EXAMPLE_SHEET, add_examples, style_sheet

VERSION = "1"
MARKER = "STT_OFFICIAL_WORKBOOK"
SCENES = {"timetable": "课表设置", "rooms": "教室设置", "school": "学校数据", "planning": "课程计划", "constraints": "约束配置"}
DOWNLOAD_SCENES = {key: SCENES[key] for key in ("rooms", "school")}
DOWNLOAD_VERSION = "3"
PROMPTS = {key: f"请按官方模板 v1 严格录入{label}。仅处理本附件填写的数据，空白可选项不覆盖原值，不删除已有数据；缺项、重名或引用不存在时停止并指出工作表和行号。生成待审批变更，不直接执行。" for key, label in SCENES.items()}
PROMPTS["timetable"] += "课表保存到本机构模板，不上传广场。"
# A star identifies a required input. Data sheets intentionally contain no sample records.
SHEETS = {
    "timetable": {"模板信息": ["模板名称*", "时间模式*", "上课日*", "显示时间轴*"],
                  "课次时段": ["模板名称*", "星期*", "节次*", "开始时间*", "结束时间*", "单元格用途*", "自定义内容"]},
    "rooms": {"教室类型": ["类型名称*", "说明"], "教室": ["教室名称*", "类型名称", "容量"]},
    "school": {"教师": ["教师姓名*", "教师分组"], "班级": ["班级名称*", "人数", "分组"], "科目": ["科目名称*", "默认课长（分钟）"]},
    "planning": {"课程计划": ["班级名称*", "科目名称*", "每周课次*", "每次课长（分钟）*", "任课教师", "固定教室"]},
    "constraints": {"规则": ["规则名称*", "规则类型*", "必须满足*", "违反扣分"],
                    "作用课次": ["规则名称*", "班级名称*", "科目名称*", "任课教师", "课次序号*"]},
}
DOWNLOAD_LAYOUTS = {
    "rooms": {"教室类型": ("教室类型", ["类型", "说明"]), "教室": ("教室清单", ["教室", "类型", "容量"])},
    "school": {"教师": ("教师清单", ["姓名", "标签"]), "班级": ("班级清单", ["班级", "人数", "分组", "班主任", "默认教室"]),
               "科目": ("科目清单", ["科目", "默认课长（分钟）"])},
}
SHEETS["school"]["班级"] += ["班主任", "默认教室"]
DOWNLOAD_COLUMN_WIDTHS = {
    "教室类型": (16, 30), "教室清单": (20, 16, 11),
    "教师清单": (14, 30), "班级清单": (16, 11, 18, 14, 22), "科目清单": (24, 24),
}
RULES = {"不重叠": "NotOverlap", "同时开始": "SameStart", "同一天": "SameDays", "不同天": "DifferentDays", "同一教室": "SameRoom"}
NOTES = {
    "timetable": ["模板信息每行一份模板；上课日用数字和逗号填写，如 1,2,3,4,5。", "课次时段每行一个单元格，星期填 1 至 7；节次从 1 开始；时间填 HH:MM，必须为 5 分钟整数倍。", "固定时段要求各上课日课次数和起止时间一致；弹性时间允许各天不同。", "支持本表声明的标准周课表结构、单元格用途及自定义文字；Excel 字体、颜色、合并、图片不作为系统模板样式导入。", "同名模板更新其完整时段。保存范围固定为本机构，不公开到广场。"],
    "rooms": ["类型可先在教室类型页填写，也可引用系统已有类型。容量为非负整数。", "类型名称必须唯一；不填写类型或容量时保留已有值。"],
    "school": ["只填写需要录入的教师、班级和科目；不填写任课关系。", "同一表内教师姓名不可重复；系统已有同名教师时需填写分组准确匹配，仍有歧义则停止处理。人数为非负整数。", "科目课长填写分钟，必须为正整数且为 5 的倍数；留空保留已有值。"],
    "planning": ["一行表示一个班级的一门科目，班级、科目、教师和教室必须已存在。", "每周课次填 1 至 50；每次课长填分钟且为 5 的倍数。", "任课教师选填，填写后同时生成授课任务；不填写时只处理课程计划。", "已有授课任务的课次数变更暂不自动增删课次，需在课程计划页核对处理。"],
    "constraints": ["规则页定义规则；作用课次页为每个规则至少填写两条不同课次。", "课次序号是对应班级、科目和任课教师任务中的第几次课；必须已存在。", "规则类型从下拉列表选择；必须满足填否时需填写非负违反扣分。", "同名规则更新；不支持自由文本规则推断，其他约束请使用普通 AI 对话。"],
}


def catalog() -> list[dict]:
    return [{"scene": key, "label": label, "version": VERSION, "prompt": PROMPTS[key],
             "sheets": [title for title, _ in DOWNLOAD_LAYOUTS[key].values()], "notes": NOTES[key],
             "download_path": f"/api/agent/official-templates/{key}.xlsx"} for key, label in DOWNLOAD_SCENES.items()]


def workbook_bytes(scene: str, sample: str = "blank") -> bytes:
    if scene == "planning" and sample == "blank":
        return planning_collection_workbook_bytes()
    if scene not in DOWNLOAD_SCENES or sample not in {"blank", "valid", "invalid", "empty"}:
        raise HTTPException(404, "没有找到此模板。")
    book = Workbook()
    book.remove(book.active)
    samples = sample_rows(scene) if sample in {"valid", "invalid"} else {}
    for title, headers in SHEETS[scene].items():
        display_title, display_headers = DOWNLOAD_LAYOUTS[scene][title]
        sheet = book.create_sheet(display_title)
        sheet.append(display_headers)
        for row in samples.get(title, []):
            sheet.append(row)
        if sample == "invalid" and title == next(iter(SHEETS[scene])):
            sheet.cell(2, 1).value = None
        _style_download_sheet(sheet, len(book.worksheets), sample)
    if sample == "blank":
        examples = sample_rows(scene)
        notes = {
            "教室类型": "类型用于给教室分类；先填类型，再在教室清单引用相同名称。",
            "教室": "每行一间教室。容量填人数数字；类型须与类型表或系统中的名称一致。可选项留空不会清除已有值。",
            "教师": "每行一位教师。姓名按实际姓名填写，标签用于区分教研组等；同名教师请补充标签并核对。",
            "班级": "每行一个行政班。人数填数字；班主任填写教师姓名，默认教室填写教室名称，须能在系统中找到。可选项可留空。",
            "科目": "每行一门科目。默认课长以分钟填写，例如40表示40分钟，需为5的倍数。",
        }
        sections = []
        for title, (display_title, headers) in DOWNLOAD_LAYOUTS[scene].items():
            rows = examples[title]
            if title == "班级":
                rows = [["一年级1班", 40, "一年级", "李老师", "一年级1班教室"], ["二年级1班", 36, "二年级", "王老师", "二年级1班教室"]]
            sections.append((display_title, headers, rows, notes[title] + " 本页仅供参考，不会导入；请在对应填写页录入真实数据。"))
        add_examples(book, sections)
    meta = book.create_sheet("_STT")
    meta.append([MARKER, DOWNLOAD_VERSION, scene])
    meta.sheet_state = "hidden"
    book.active = 0
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def planning_collection_workbook_bytes() -> bytes:
    """Lesson-per-row collection layout; use the AI parser, not the legacy v1 contract."""
    book = Workbook()
    sheet = book.active
    sheet.title = "课程计划"
    sheet.append(["班级", "科目", "教师", "默认教室", "课次名", "课次教室"])
    edge = Side(style="thin", color="BFBFBF")
    for row in sheet.iter_rows(min_row=1, max_row=14, max_col=6):
        for cell in row:
            cell.font = Font(name="等线", size=12, bold=cell.row == 1, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="D9E1F2")
        sheet.row_dimensions[row[0].row].height = 24
    # Match the reference's class/course grouping; each unmerged lesson row is one lesson.
    sheet.merge_cells("A2:A14")
    for start, end in ((2, 5), (6, 9), (10, 11), (12, 14)):
        for column in "BCD":
            sheet.merge_cells(f"{column}{start}:{column}{end}")
    for column, width in zip("ABCDEF", (18, 18, 16, 24, 20, 24)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.print_options.horizontalCentered = True
    sheet.print_area = "A1:F14"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    style_sheet(sheet, (18, 18, 16, 24, 20, 24))
    example_sheet = add_examples(book, [("课程计划：每行一节课", ["班级", "科目", "教师", "默认教室", "课次名", "课次教室"], [
        ["一年级1班", "语文", "李老师", "一年级1班教室", "语文第1课次", "默认"],
        ["一年级1班", "语文", "李老师", "一年级1班教室", "语文第2课次", "默认"],
        ["一年级1班", "科学", "王老师", "一年级1班教室", "科学实验", "科学一室"],
        ["一年级1班", "体育", "张老师", "不使用教室", "体育第1课次", "不使用教室"],
    ], "每个填写了课次名的行是一节课，例中语文每周2节。合并单元格覆盖的行共用班级、科目、教师和默认教室；可按实际课次数取消或调整合并，也可每行重复填写。课次教室填“默认”时沿用默认教室；填具体名称时改用该教室。课长沿用科目默认课长。本页不会导入。")])
    for column in "ABCD":
        example_sheet.merge_cells(f"{column}3:{column}4")
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def _style_download_sheet(sheet, index: int, sample: str) -> None:
    style_sheet(sheet, DOWNLOAD_COLUMN_WIDTHS[sheet.title])


def sample_rows(scene: str) -> dict[str, list[list]]:
    return {
        "rooms": {"教室类型": [["科学实验室", "科学教学"]], "教室": [["科学一室", "科学实验室", 40], ["科学二室", "科学实验室", 32]]},
        "school": {"教师": [["李老师", "语文组"], ["王老师", "数学组"]], "班级": [["一年级1班", 40, "一年级"], ["二年级1班", 36, "二年级"]], "科目": [["语文", 40], ["数学", 40]]},
    }[scene]


def parse_official_workbook(data: bytes) -> dict | None:
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 64 * 1024 * 1024:
                raise HTTPException(413, "Excel 解压后超过 64 MB，请拆分文件。")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except (BadZipFile, ValueError, KeyError) as exc:
        raise HTTPException(400, "无法读取 Excel 文件，请重新另存为 .xlsx。") from exc
    try:
        if "_STT" not in book.sheetnames:
            return None
        marker, version, scene = [book["_STT"].cell(1, col).value for col in (1, 2, 3)]
        if marker != MARKER or str(version) not in {VERSION, "2", DOWNLOAD_VERSION} or scene not in SCENES or (str(version) == "2" and scene not in LAYOUTS) or (str(version) == DOWNLOAD_VERSION and scene not in DOWNLOAD_LAYOUTS):
            raise HTTPException(400, "官方模板标识或版本无效，请重新下载模板。")
        if str(version) == "2":
            rows = read_layout(book, scene)
            return {"kind": "stt_official_workbook", "version": VERSION, "workbook_version": "2", "scene": scene,
                    "sheets": [{"name": name, "headers": SHEETS[scene][name], "rows": values} for name, values in rows.items()]}
        layout = DOWNLOAD_LAYOUTS[scene] if str(version) == DOWNLOAD_VERSION else {title: (title, headers) for title, headers in SHEETS[scene].items()}
        expected_sheets = {"_STT", *(title for title, _ in layout.values())}
        actual_sheets = set(book.sheetnames) - {EXAMPLE_SHEET}
        if actual_sheets != expected_sheets | {"填写说明"} and not (scene in DOWNLOAD_SCENES and actual_sheets == expected_sheets):
            raise HTTPException(400, "工作表名称不匹配，请勿新增、删除或重命名官方模板工作表。")
        result = {"kind": "stt_official_workbook", "version": VERSION, "scene": scene, "sheets": []}
        for title, headers in SHEETS[scene].items():
            display_title, display_headers = layout[title]
            sheet = book[display_title]
            if sheet.max_row > 2001 or sheet.max_column > len(headers):
                raise HTTPException(400, f"{title}：最多填写 2000 条数据，且不能添加列。")
            actual_headers = [sheet.cell(1, i).value for i in range(1, len(headers) + 1)]
            legacy_class_headers = scene == "school" and title == "班级" and str(version) == VERSION and actual_headers == headers[:3] + [None, None]
            if actual_headers != display_headers and not legacy_class_headers:
                raise HTTPException(400, f"{title}：表头不匹配，请使用原始表头。")
            rows = []
            for index, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(headers)), 2):
                if not any(cell.value is not None and str(cell.value).strip() for cell in cells):
                    continue
                if any(cell.data_type == "f" for cell in cells):
                    raise HTTPException(400, f"{title}第 {index} 行：不支持公式，请粘贴为值。")
                values = [cell.value.strftime("%H:%M") if isinstance(cell.value, time) else cell.value for cell in cells]
                if any(len(str(value or "")) > 2000 for value in values):
                    raise HTTPException(400, f"{title}第 {index} 行：单元格内容过长。")
                rows.append({"row": index, "values": values, "source_sheet": display_title})
            result["sheets"].append({"name": title, "headers": headers, "rows": rows})
        return result
    finally:
        book.close()


def official_data(attachment: dict) -> dict | None:
    data = attachment.get("extracted_data") or {}
    return data if data.get("kind") == "stt_official_workbook" else None


def proposals_from_workbook(attachment: dict, scene: str, repo, organization_id: str, project_id: str | None) -> list[dict]:
    data = official_data(attachment) or {}
    if data.get("scene") != scene or data.get("version") != VERSION:
        raise ActionCompilationError([{"code": "official_scene_mismatch", "message": "附件场景与当前会话不一致，请切换到对应场景。"}])
    proposals, rows_by_sheet = [], {}
    locations = {(sheet["name"], raw["row"]): raw.get("source_sheet", sheet["name"])
                 for sheet in data.get("sheets", []) for raw in sheet["rows"]}

    def fail(sheet, row, message):
        sheet = locations.get((sheet, row), sheet)
        raise ActionCompilationError([{"code": "official_validation", "message": f"{sheet}第 {row} 行：{message}", "sheet": sheet, "row": row}])

    def number(value, sheet, row, label, low=0, high=100000):
        try:
            parsed = float(value)
            if not parsed.is_integer() or not low <= parsed <= high:
                raise ValueError()
            return int(parsed)
        except (ValueError, TypeError, OverflowError):
            fail(sheet, row, f"{label}须为 {low} 至 {high} 的整数。")

    def minutes(value, sheet, row):
        parsed = number(value, sheet, row, "课长", 5, 1440)
        if parsed % 5:
            fail(sheet, row, "课长须为 5 分钟的倍数。")
        return parsed // 5

    def choice(value, choices, sheet, row, label):
        if value not in choices:
            fail(sheet, row, f"{label}只能填写：{'、'.join(choices)}。")
        return choices[value]

    for sheet in data.get("sheets", []):
        title = sheet["name"]
        if title not in SHEETS[scene]:
            fail(title, 1, "未知工作表。")
        rows_by_sheet[title] = []
        for raw in sheet["rows"]:
            row = raw["row"]
            values = {header.rstrip("*"): str(value).strip() if isinstance(value, str) else value for header, value in zip(SHEETS[scene][title], raw["values"])}
            for header in SHEETS[scene][title]:
                if header.endswith("*") and values.get(header.rstrip("*")) in (None, ""):
                    fail(title, row, f"请填写{header.rstrip('*')}。")
            rows_by_sheet[title].append((row, values))
    if not any(rows_by_sheet.values()):
        fail(next(iter(SHEETS[scene])), 1, "模板尚未填写，没有可录入的数据。")
    school = repo.list_school_data(organization_id)
    planning = repo.list_planning_data(organization_id, project_id) if project_id else {}
    if scene == "constraints" and project_id:
        planning = {**planning, "constraints": repo.list_constraints(organization_id, project_id).get("constraints", [])}
    seen = set()

    def unique(target, key, sheet, row):
        identity = (target, str(key))
        if identity in seen:
            fail(sheet, row, "存在重复记录，请合并后再提交。")
        seen.add(identity)

    def reference(target, name, sheet, row):
        collection = {"teacher": "teachers", "homeroom": "homerooms", "subject": "subjects", "room": "rooms", "room_type": "room_types"}[target]
        matches = [item for item in school.get(collection, []) if item.get("name") == name]
        if len(matches) != 1:
            fail(sheet, row, f"“{name}”不存在或存在重名，请先核对基础数据。")
        return str(matches[0]["id"])

    def add(target, payload, sheet, row, operation="bulk_upsert"):
        sheet = locations.get((sheet, row), sheet)
        payload["source_ref"] = {"filename": attachment.get("filename"), "sheet": sheet, "row": row}
        proposals.append({"target": target, "operation": operation, "payload": {"items": [payload]} if operation == "bulk_upsert" else payload, "human_summary": f"按官方模板录入{SCENES[scene]}"})

    if scene in {"school", "rooms"}:
        mapping = {"教师": ("teacher", "教师姓名", {"教师分组": "department"}), "班级": ("homeroom", "班级名称", {"人数": "student_count", "分组": "group_name", "班主任": "teacher_name", "默认教室": "room_name"}),
                   "科目": ("subject", "科目名称", {"默认课长（分钟）": "default_duration_slots"}), "教室类型": ("room_type", "类型名称", {"说明": "description"}), "教室": ("room", "教室名称", {"类型名称": "room_type_name", "容量": "capacity"})}
        for title in SHEETS[scene]:
            target, name_key, fields = mapping[title]
            for row, values in rows_by_sheet.get(title, []):
                unique(target, values[name_key], title, row)
                payload = {"name": values[name_key]}
                for label, field in fields.items():
                    value = values.get(label)
                    if value in (None, ""):
                        continue
                    payload[field] = minutes(value, title, row) if field == "default_duration_slots" else number(value, title, row, label) if field in {"student_count", "capacity"} else value
                add(target, payload, title, row)
    elif scene == "planning":
        for row, values in rows_by_sheet.get("课程计划", []):
            title = "课程计划"
            homeroom = reference("homeroom", values["班级名称"], title, row)
            subject = reference("subject", values["科目名称"], title, row)
            unique("course_plan", (homeroom, subject), title, row)
            payload = {"homeroom_id": homeroom, "subject_id": subject, "weekly_slots": number(values["每周课次"], title, row, "每周课次", 1, 50), "duration_slots": minutes(values["每次课长（分钟）"], title, row)}
            add("course_plan", dict(payload), title, row)
            if values.get("任课教师"):
                payload["primary_teacher_id"] = reference("teacher", values["任课教师"], title, row)
                if values.get("固定教室"):
                    payload["fixed_room_id"] = reference("room", values["固定教室"], title, row)
                tasks = [task for task in planning.get("teaching_tasks", []) if task.get("homeroom_id") == homeroom and task.get("subject_id") == subject]
                if tasks:
                    if len(tasks) != 1 or tasks[0].get("primary_teacher_id") not in (None, "", payload["primary_teacher_id"]):
                        fail(title, row, "该课程存在不同或多条授课任务，请在课程计划页面核对。")
                    lessons = [lesson for lesson in planning.get("task_lessons", []) if lesson.get("teaching_task_id") == tasks[0]["id"]]
                    if len(lessons) != payload["weekly_slots"] or any(int(lesson.get("duration_slots") or 0) != payload["duration_slots"] for lesson in lessons):
                        fail(title, row, "已有任务的课次数或课长不同，请先在课程计划页面调整。")
                add("task", payload, title, row)
            elif values.get("固定教室"):
                fail(title, row, "填写固定教室时也需要填写任课教师。")
    elif scene == "timetable":
        templates = {}
        for row, values in rows_by_sheet.get("模板信息", []):
            title, name = "模板信息", values["模板名称"]
            unique("template", name, title, row)
            mode = choice(values["时间模式"], {"固定时段": "fixed", "弹性时间": "variable"}, title, row, "时间模式")
            days = str(values["上课日"]).replace("，", ",").split(",")
            days = [number(day.strip(), title, row, "上课日", 1, 7) for day in days]
            if len(set(days)) != len(days):
                fail(title, row, "上课日不能重复。")
            templates[name] = ({"name": name, "day_count": 7, "slot_duration_minutes": 5, "template_kind": "normal", "application_scope": "all", "preset_scope": "organization", "publish_as_preset": False,
                "display_config": {"title": name, "enabled_weekdays": sorted(days), "time_mode": mode, "show_time_axis": choice(values["显示时间轴"], {"是": True, "否": False}, title, row, "显示时间轴"), "preset_scope": "organization", "course_cell_layout": "two_line", "course_cell_top_field": "subject", "course_cell_bottom_field": "teacher"}, "periods": []}, row)
        for row, values in rows_by_sheet.get("课次时段", []):
            title, name = "课次时段", values["模板名称"]
            if name not in templates:
                fail(title, row, "模板名称未在模板信息页定义。")
            payload, _ = templates[name]
            day = number(values["星期"], title, row, "星期", 1, 7)
            ordinal = number(values["节次"], title, row, "节次", 1, 50)
            unique("period", (name, day, ordinal), title, row)
            if day not in payload["display_config"]["enabled_weekdays"]:
                fail(title, row, "该星期未在模板信息中启用。")
            times = []
            for label in ("开始时间", "结束时间"):
                value = str(values[label])
                if not re.fullmatch(r"(?:[01]?\d|2[0-3]):[0-5]\d", value):
                    fail(title, row, f"{label}请使用 HH:MM 格式。")
                hour, minute = map(int, value.split(":"))
                if minute % 5:
                    fail(title, row, "时间必须为 5 分钟的整数倍。")
                times.append(hour * 60 + minute)
            if times[0] >= times[1]:
                fail(title, row, "结束时间必须晚于开始时间。")
            mode = choice(values["单元格用途"], {"安排课次": "scheduled", "自定义内容": "custom", "不安排": "disabled"}, title, row, "单元格用途")
            content = values.get("自定义内容") or ""
            if mode == "custom" and not content or mode != "custom" and content:
                fail(title, row, "自定义内容只能且必须在“自定义内容”用途时填写。")
            payload["periods"].append({"weekday": day, "period_index": ordinal, "label": f"第{ordinal}节", "start_time_minutes": times[0], "end_time_minutes": times[1], "active": mode == "scheduled", "display_config": {"cell_mode": mode, "custom_content": content, "align": "center", "colspan": 1}})
        for payload, row in templates.values():
            schedules = []
            for day in payload["display_config"]["enabled_weekdays"]:
                periods = sorted([p for p in payload["periods"] if p["weekday"] == day], key=lambda p: p["period_index"])
                if not periods or [p["period_index"] for p in periods] != list(range(1, len(periods) + 1)):
                    fail("模板信息", row, f"星期 {day} 的课次须从 1 连续填写。")
                if any(a["end_time_minutes"] > b["start_time_minutes"] for a, b in zip(periods, periods[1:])):
                    fail("模板信息", row, f"星期 {day} 的课次时间存在重叠或倒序。")
                schedules.append([(p["start_time_minutes"], p["end_time_minutes"]) for p in periods])
            if payload["display_config"]["time_mode"] == "fixed" and any(s != schedules[0] for s in schedules):
                fail("模板信息", row, "固定时段要求各天时间一致；请修正或选择弹性时间。")
            add("timetable_template", payload, "模板信息", row)
    elif scene == "constraints":
        groups = {}
        for row, values in rows_by_sheet.get("作用课次", []):
            title = "作用课次"
            home = reference("homeroom", values["班级名称"], title, row)
            subject = reference("subject", values["科目名称"], title, row)
            teacher = reference("teacher", values["任课教师"], title, row) if values.get("任课教师") else None
            tasks = [task for task in planning.get("teaching_tasks", []) if task.get("homeroom_id") == home and task.get("subject_id") == subject and (not teacher or task.get("primary_teacher_id") == teacher)]
            if len(tasks) != 1:
                fail(title, row, "授课任务不存在或不唯一，请补全任课教师并核对课程计划。")
            all_lessons = values["课次序号"] == "全部"
            ordinal = None if all_lessons else number(values["课次序号"], title, row, "课次序号", 1, 50)
            matches = [lesson for lesson in planning.get("task_lessons", []) if lesson.get("teaching_task_id") == tasks[0]["id"] and (all_lessons or int(lesson.get("lesson_index") or 0) == ordinal)]
            if not matches or (not all_lessons and len(matches) != 1):
                fail(title, row, "找不到指定课次。")
            for lesson in matches:
                unique("rule_lesson", (values["规则名称"], lesson["id"]), title, row)
                groups.setdefault(values["规则名称"], []).append(str(lesson["id"]))
        for row, values in rows_by_sheet.get("规则", []):
            title, name = "规则", values["规则名称"]
            unique("constraint", name, title, row)
            lesson_ids = groups.pop(name, [])
            if len(lesson_ids) < 2:
                fail(title, row, "每个规则至少需要两条不同的作用课次。")
            required = choice(values["必须满足"], {"是": True, "否": False}, title, row, "必须满足")
            penalty = number(values.get("违反扣分"), title, row, "违反扣分") if not required or values.get("违反扣分") not in (None, "") else None
            payload = {"name": name, "distribution_type": choice(values["规则类型"], RULES, title, row, "规则类型"), "required": required, "penalty": penalty, "parameters": {"lesson_ids": lesson_ids}}
            existing = [c for c in planning.get("constraints", []) if c.get("name") == name]
            if len(existing) > 1:
                fail(title, row, "系统存在多条同名规则。")
            if existing:
                payload["id"] = existing[0]["id"]
            add("constraint", payload, title, row, "update" if existing else "create")
        if groups:
            fail("作用课次", 2, "存在未在规则页定义的规则名称。")
    return proposals
