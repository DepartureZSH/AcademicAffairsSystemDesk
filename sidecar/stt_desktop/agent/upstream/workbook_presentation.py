"""Shared presentation for downloadable data collection workbooks."""
from math import ceil
from unicodedata import east_asian_width
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EXAMPLE_SHEET = "填写示例（不导入）"


def style_sheet(sheet, widths, *, last_row=14, row_height=28, header_row=1):
    edge = Side(style="thin", color="BFC9D9")
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=header_row, max_row=max(last_row, sheet.max_row), max_col=len(widths)):
        for cell in row:
            heading = cell.row == header_row
            cell.font = Font(name="等线", size=12, bold=heading, color="243247")
            cell.fill = PatternFill("solid", fgColor="D9E5F5" if heading else "FFFFFF")
            cell.alignment = Alignment(horizontal="center" if heading else "left", vertical="center", wrap_text=True)
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
        lines = max(ceil(sum(2 if east_asian_width(char) in "WF" else 1 for char in str(cell.value or "")) / max(width - 2, 1)) for cell, width in zip(row, widths))
        sheet.row_dimensions[row[0].row].height = max(32 if row[0].row == header_row else row_height, lines * 18 + 10)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.sheet_view.showGridLines = False


def add_examples(book, sections):
    sheet = book.create_sheet(EXAMPLE_SHEET)
    for title, headers, rows, explanation in sections:
        start = sheet.max_row + 2 if sheet.cell(1, 1).value is not None else 1
        sheet.cell(start, 1, title)
        sheet.merge_cells(start_row=start, start_column=1, end_row=start, end_column=len(headers))
        sheet.cell(start, 1).font = Font(name="等线", size=12, bold=True, color="243247")
        sheet.row_dimensions[start].height = 30
        for col, label in enumerate(headers, 1):
            sheet.cell(start + 1, col, label)
        for index, values in enumerate(rows, start + 2):
            for col, value in enumerate(values, 1):
                sheet.cell(index, col, value)
        end = start + 1 + len(rows)
        style_sheet(sheet, [26] * len(headers), header_row=start + 1, last_row=end, row_height=60)
        sheet.cell(end + 1, 1, explanation)
        sheet.merge_cells(start_row=end + 1, start_column=1, end_row=end + 1, end_column=len(headers))
        sheet.cell(end + 1, 1).alignment = Alignment(vertical="center", wrap_text=True)
        sheet.cell(end + 1, 1).font = Font(name="等线", size=12, color="243247")
        sheet.row_dimensions[end + 1].height = 80
    sheet.freeze_panes = None
    return sheet
