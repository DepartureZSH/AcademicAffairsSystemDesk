"""Six-column constraint workbooks. Notes remain local review data, never AI input."""
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from stt_desktop.agent.upstream.workbook_presentation import EXAMPLE_SHEET, add_examples, style_sheet


class WorkbookRowAnalysis(BaseModel):
    query: str = Field(min_length=1, max_length=1000)
    scope: str = Field(min_length=1, max_length=1000)
    required: bool = True
    penalty: int | None = Field(default=10)


def analyze_row(service, user_id, organization_id, project_id, payload):
    from stt_desktop.agent.upstream.constraint_wizard import search_rules
    from stt_desktop.agent.upstream.constraint_resolver import validate_strength
    penalty = validate_strength(payload["required"], payload.get("penalty"))
    found = search_rules(service, user_id, organization_id, project_id, payload["query"], False)
    result = {"rules": found["items"], "message": found["message"], "resolution": None}
    if len(found["items"]) == 1 and found["items"][0]["match_label"] == "高匹配":
        rule = found["items"][0]
        result["resolution"] = service.resolve_constraint_scope(user_id, organization_id, project_id, {
            "template_id": rule["template_id"], "scenario_type": rule["scenario_type"], "constraint_description": rule["name"],
            "required": payload["required"], "penalty": penalty, "scope_text": payload["scope"], "preview_only": True,
        })
    return result

HEADERS = ["约束规则", "约束强度", "涉及课次", "是否相似扩展", "相似扩展规则", "备注"]
TEMPLATES = [
    {"id": "blank", "name": "空白约束表", "rule": "", "note": "每行一条约束；强度填写必须满足或尽量满足（10/30/60）；涉及课次请写明班级、科目、教师及第几课次。"},
    {"id": "parallel", "name": "场地同时上课上限", "rule": "同一时段最多同时上4节课", "note": "例如涉及课次填全校体育课。所选课次合为一组共享4个名额，不是每班各4个。按实际场地容量修改规则。"},
    {"id": "consecutive", "name": "两节课连续安排", "rule": "两节课按顺序连续安排", "note": "例如一年级1班语文张老师第1、2课次。相似扩展可写换到其他班级的语文课，保持第1、2课次。"},
    {"id": "spread", "name": "课程分散到不同天", "rule": "所选课次分别安排在不同星期几", "note": "例如一年级1班体育第1、2、3课次。涉及课次不要超过可上课天数；每个授课任务分别创建。"},
]


def build_workbook(template_id: str) -> bytes:
    template = next((t for t in TEMPLATES if t["id"] == template_id), None)
    if not template:
        raise ValueError("没有找到该模板。")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "约束填写表"
    sheet.append(HEADERS)
    sheet.append(["", "", "", "", "", ""])
    widths = [40, 24, 46, 20, 46, 66]
    style_sheet(sheet, widths, last_row=12, row_height=72)
    samples = {
        "blank": ["张老师的课不要同时上", "必须满足", "张老师的全部课次", "否", "", "先写清要遵守的规则，再写哪些课参与。"],
        "parallel": [template["rule"], "必须满足", "全校体育课的全部课次", "否", "", template["note"]],
        "consecutive": [template["rule"], "尽量满足（30）", "一年级1班语文李老师第1、2课次", "是", "扩展到一年级其他班级的语文第1、2课次，各班分别创建", template["note"]],
        "spread": [template["rule"], "必须满足", "一年级1班体育张老师第1、2、3课次", "是", "扩展到其他班级的体育课，每班分别创建，不跨班混合", template["note"]],
    }
    add_examples(workbook, [(template["name"], HEADERS, [samples[template_id]],
        "每行一条约束。强度填“必须满足”或“尽量满足（10/30/60）”，数字越大越优先。涉及课次写清班级、科目、教师和第几课次。扩展填“是”时，必须写明哪些条件变化、哪些保持不变；不扩展填“否”，扩展规则留空。备注只供人阅读，不交给AI。示例不会导入，请在约束填写表录入真实要求。")])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def read_workbook(data: bytes) -> list[dict]:
    if len(data) > 2 * 1024 * 1024:
        raise ValueError("文件不能超过 2 MB。")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.infolist()) > 200 or sum(i.file_size for i in archive.infolist()) > 16 * 1024 * 1024:
                raise ValueError("文件解压后过大，请使用小型约束填写表。")
            if any("externalLinks/" in i.filename for i in archive.infolist()):
                raise ValueError("模板不支持外部链接，请移除外部引用。")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("无法读取文件，请上传有效的 .xlsx 文件。") from exc
    try:
        data_sheets = [sheet for sheet in workbook.worksheets if sheet.title != EXAMPLE_SHEET]
        if len(data_sheets) != 1:
            raise ValueError("请只保留一个约束填写表工作表。")
        sheet = data_sheets[0]
        if sheet.max_column > 6 or sheet.max_row > 101:
            raise ValueError("模板需为六列，最多 100 个填写行，每批最多解析 20 条约束。")
        rows = list(sheet.iter_rows())
        if not rows or [str(c.value or "").strip() for c in rows[0]] != HEADERS:
            raise ValueError("表头应依次为：" + "、".join(HEADERS))
        result = []
        for index, cells in enumerate(rows[1:], 2):
            if any(c.data_type == "f" for c in cells):
                raise ValueError(f"第 {index} 行包含公式，请改为填写文字。")
            values = [str(c.value or "").strip() for c in cells]
            if not any(values[:5]):
                continue
            if any(len(v) > 1000 for v in values):
                raise ValueError(f"第 {index} 行文字过长，每格最多 1000 字。")
            rule, strength, scope, expand, instruction, notes = values
            error = ""
            if not rule or not scope:
                error = "请填写约束规则和涉及课次。"
            required = strength in ("必须满足", "硬约束")
            penalty = None
            if not required:
                normalized = strength.replace("（", "(").replace("）", ")").replace(" ", "")
                options = {"尽量满足": 10, "尽量满足(10)": 10, "尽量满足(30)": 30, "尽量满足(60)": 60}
                penalty = options.get(normalized)
                if penalty is None:
                    error = "强度请填写必须满足或尽量满足（10/30/60）。"
            if expand not in ("是", "否", ""):
                error = "是否相似扩展请填写是或否。"
            if expand == "是" and not instruction:
                error = "请填写相似扩展规则，避免 AI 猜测扩展范围。"
            result.append({"row": index, "query": rule, "scope": scope, "required": required,
                           "penalty": penalty or 10, "expand": expand == "是", "instruction": instruction,
                           "notes": notes, "error": error})
        if len(result) > 20:
            raise ValueError("每批最多 20 条约束，请分批上传。")
        if not result:
            raise ValueError("表格中还没有填写约束。")
        return result
    finally:
        workbook.close()
